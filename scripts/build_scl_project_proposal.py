from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SCL_Project_Proposal.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(89, 89, 89)
LIGHT_FILL = "F4F6F9"
BORDER = "D7DBE2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    set_table_borders(table)


def mark_first_row_header(table):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph(doc, text="", style=None, after=8, before=0, align=None, bold=False, italic=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, bold=bold, italic=italic, color=color)
    return p


def bullet(doc, text):
    p = paragraph(doc, style="List Bullet", after=4)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.add_run(text)
    return p


def numbered(doc, text):
    p = paragraph(doc, style="List Number", after=4)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(10 if level == 1 else 6 if level == 2 else 4)
    for run in p.runs:
        set_font(run, size=16 if level == 1 else 13 if level == 2 else 12, color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_source(doc, label, url):
    p = paragraph(doc, after=3)
    r = p.add_run(label + ": ")
    set_font(r, bold=True)
    p.add_run(url)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [6.5])
    mark_first_row_header(table)
    set_cell_shading(table.cell(0, 0), LIGHT_FILL)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, bold=True, color=NAVY)
    p.add_run("\n" + body)
    paragraph(doc, after=4)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_widths(table, widths)
    mark_first_row_header(table)
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_font(r, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(val)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_widths(table, widths)
    paragraph(doc, after=6)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color in (("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = "Sports Cappers Leaderboard | Modernization Proposal"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = "Prepared draft - June 18, 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_font(run, size=9, color=GRAY)

    # Proposal centerpiece title block.
    paragraph(doc, "Project Proposal", after=8, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=GRAY)
    title = paragraph(doc, "Sports Cappers Leaderboard Modernization", after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in title.runs:
        set_font(run, size=24, color=RGBColor(0, 0, 0), bold=True)
    subtitle = paragraph(doc, "Hybrid Winible Affiliate + SCL Direct Marketplace Rebuild", after=8, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)
    for run in subtitle.runs:
        set_font(run, size=14, color=GRAY)
    paragraph(doc, "Prepared for the owners of Sports Cappers Leaderboard", after=22, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=GRAY)

    meta = doc.add_table(rows=4, cols=2)
    set_table_widths(meta, [1.7, 4.8])
    mark_first_row_header(meta)
    rows = [
        ("Prepared date", "June 18, 2026"),
        ("Primary objective", "Modernize SCL into a trusted marketplace for verified handicappers and betting picks."),
        ("Commerce model", "Keep Winible affiliate package revenue active while building a foundation for SCL Direct subscriptions."),
        ("Recommended delivery", "Phased rebuild, starting with accounts, leaderboard upgrades, package governance, and improved play entry."),
    ]
    for r_idx, (label, value) in enumerate(rows):
        label_cell, value_cell = meta.rows[r_idx].cells
        set_cell_shading(label_cell, LIGHT_FILL)
        p = label_cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), bold=True, color=NAVY)
        value_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        value_cell.paragraphs[0].add_run(value)
    paragraph(doc, after=10)

    add_callout(
        doc,
        "Recommended Positioning",
        "SCL should be rebuilt as the trust and comparison layer for sports handicappers: verified records, transparent performance, structured plays, and package discovery. The first release should protect the current Winible affiliate model while introducing the account, package, and data architecture needed for SCL Direct commerce."
    )

    heading(doc, "Executive Summary")
    paragraph(doc, "Sports Cappers Leaderboard already has a strong core premise: users can compare handicappers by record, units, ROI, sport, and time period. The current site, however, reads more like an older directory than a modern subscription marketplace. The next version should turn SCL into a product users can trust, revisit, and transact through.")
    paragraph(doc, "The proposal is to rebuild SCL around a hybrid package system. Some cappers continue selling through Winible, with SCL-controlled affiliate package links. Other cappers, when approved, can sell directly through SCL using Stripe. To the customer, packages should feel consistent across the site; behind the scenes, each package has a defined checkout model, attribution policy, access-control behavior, and admin approval trail.")
    paragraph(doc, "Phase 1 should focus on the platform foundation: modern redesign, user accounts, capper accounts, structured play entry, improved profiles, robust leaderboard filtering, and admin-governed Winible package links. Phase 2 adds SCL Direct checkout, subscriptions, premium play access, and subscription management. Phase 3 expands into deeper analytics, revenue reporting, capper payouts, closing-line-value tracking, and automation.")

    heading(doc, "Current-State Observations")
    bullet(doc, "SCL already presents verified performance themes, including rankings by sport, rankings by time period, records, win percentage, units, and ROI.")
    bullet(doc, "The public site includes capper profile/store destinations and calls to action for VIP trials, but it does not appear to offer regular user accounts, customer dashboards, or native subscription management.")
    bullet(doc, "Winible is currently important to SCL's revenue model because capper packages can be sold through Winible storefront/package links while SCL maintains affiliate attribution.")
    bullet(doc, "Winible's creator tooling is broader than simple checkout: its help materials describe storefronts, pick posting, paid/free picks, subscription plans, billing variants, trials, promo codes, customer management, earnings, and Stripe-based payouts.")
    bullet(doc, "The rebuild should avoid giving cappers unrestricted control over package links, because package-level affiliate attribution is central to current revenue protection.")

    heading(doc, "Project Goals")
    numbered(doc, "Modernize the public SCL experience so users can discover, compare, follow, and evaluate cappers with more confidence.")
    numbered(doc, "Preserve the existing Winible affiliate revenue model during and after the redesign.")
    numbered(doc, "Create the account and data foundation for SCL Direct subscriptions without forcing that model into Phase 1.")
    numbered(doc, "Make play entry structured, sportsbook-like, and useful for tracking performance instead of relying on loose text fields.")
    numbered(doc, "Give SCL admins durable control over capper approval, package approval, affiliate links, leaderboard settings, content moderation, and revenue visibility.")

    heading(doc, "Recommended Solution")
    paragraph(doc, "The rebuild should separate marketplace presentation from commerce execution. Public users see one coherent SCL marketplace, while the system determines whether each package checks out through a Winible affiliate URL or through SCL Direct.")
    add_matrix(
        doc,
        ["Component", "Recommended Build"],
        [
            ("Public site", "Modern homepage, leaderboard, capper profiles, package listings, today/yesterday picks, sport pages, and trust-focused comparison surfaces."),
            ("User accounts", "Secure signup/login, followed cappers, subscriptions, accessible plays, profile settings, and billing links where applicable."),
            ("Capper accounts", "Profile management, play entry, performance view, package requests, and status visibility for SCL-approved store listings."),
            ("Admin dashboard", "Capper approval, package approval, final URL control, content moderation, leaderboard settings, user oversight, analytics, and commerce status."),
            ("Data model", "Structured plays, normalized bet types, package types, account roles, permissions, audit events, and source-controlled affiliate links."),
        ],
        [1.7, 4.8],
    )

    heading(doc, "Hybrid Package Model")
    paragraph(doc, "Every package in SCL should have a package type. This design prevents the current affiliate model and the future direct subscription model from competing in the database, admin workflow, and user interface.")
    add_matrix(
        doc,
        ["Package Type", "Checkout Destination", "Revenue Handling", "Admin Control"],
        [
            ("Winible Package", "Approved Winible package URL with SCL affiliate attribution.", "Revenue tracked through Winible affiliate arrangement.", "SCL owns the final displayed URL and can approve, reject, edit, disable, or replace links."),
            ("SCL Direct Package", "Stripe Checkout or equivalent SCL-hosted checkout flow.", "SCL collects payment, manages access, and later shares revenue with capper based on agreed terms.", "SCL controls package availability, pricing rules, subscriptions, refunds policy, and access state."),
        ],
        [1.45, 1.85, 1.7, 1.5],
    )
    paragraph(doc, "Capper package changes should use a request-and-approval workflow. A capper can request a new package or edits to an existing package, but the package does not become public until an admin reviews the details and confirms the checkout URL. This protects package-level attribution and keeps the marketplace consistent.")

    heading(doc, "Structured Play Entry")
    paragraph(doc, "The play-entry system should feel closer to a sportsbook or professional bet-tracking platform than a form for free-text picks. Cappers should enter structured data so leaderboard logic, profile analytics, and premium access can rely on clean records.")
    add_matrix(
        doc,
        ["Field Group", "Examples"],
        [
            ("Event", "Sport, league, game/event, start time, team/player."),
            ("Bet", "Moneyline, spread, total, player prop, team prop, future, parlay, DFS, other."),
            ("Line", "Sportsbook, odds, line, units, selection, confidence rating."),
            ("Access", "Free/premium designation, package/category, subscriber visibility."),
            ("Tracking", "Pending, win, loss, push, void, closing line, optional CLV."),
        ],
        [1.5, 5.0],
    )

    heading(doc, "Phase Roadmap")
    add_matrix(
        doc,
        ["Phase", "Scope", "Outcome"],
        [
            ("Phase 1: Marketplace Foundation", "Modern responsive redesign, user accounts, capper accounts, improved profiles, better leaderboard filters, structured play entry, admin package approval, and Winible affiliate links.", "SCL looks modern, protects existing revenue, and has the database foundation for future direct sales."),
            ("Phase 2: SCL Direct Commerce", "Stripe Checkout, recurring subscriptions, package tiers, trials/promo codes if desired, premium content access, billing portal links, and admin subscription/revenue views.", "Customers can subscribe directly on SCL while Winible packages continue to operate in parallel."),
            ("Phase 3: Advanced Analytics", "Capper revenue share dashboard, deeper charts, sport splits, CLV, automation, notifications, and payout reporting.", "SCL becomes a higher-trust analytics and commerce platform rather than only a leaderboard."),
        ],
        [1.55, 3.25, 1.7],
    )

    heading(doc, "Security, Privacy, and Compliance Requirements")
    bullet(doc, "Passwords must be hashed and never visible to admins.")
    bullet(doc, "Authentication should include secure login, logout, password reset, rate limiting, and role-based access for admin, capper, and user roles.")
    bullet(doc, "Collect only the personal information needed for accounts, subscriptions, support, security, and compliance.")
    bullet(doc, "Use HTTPS everywhere, secure cookies, CSRF protection where applicable, audit logs for admin actions, and least-privilege permissions.")
    bullet(doc, "Include privacy policy, cookie notice/consent where tracking pixels are used, refund policy, terms, disclaimer, and responsible-gaming messaging.")
    bullet(doc, "For SCL Direct, use hosted Stripe surfaces where practical to reduce payment-data handling and accelerate PCI scope management.")

    heading(doc, "Stripe Approach for SCL Direct")
    paragraph(doc, "For Phase 2, the most practical first version is Stripe Checkout plus Stripe Billing for recurring packages, with Stripe's customer portal for billing management. If SCL later wants automated capper payouts or platform fees, Stripe Connect can support marketplace-style subscriptions and application fees. The final charge architecture should be confirmed during technical discovery because it affects merchant-of-record responsibility, tax handling, payout timing, refunds, and support obligations.")

    heading(doc, "Proposed Deliverables")
    bullet(doc, "Product discovery and technical specification, including data model, role model, package model, and admin approval workflows.")
    bullet(doc, "Modern responsive UI/UX for homepage, leaderboard, capper profile, package/store views, auth flows, and dashboards.")
    bullet(doc, "Backend application with user, capper, admin, play, package, leaderboard, and audit-log modules.")
    bullet(doc, "Winible package-link governance with admin-owned affiliate URLs and capper request workflow.")
    bullet(doc, "Structured play-entry system and leaderboard calculations for records, units, ROI, win percentage, streaks, and sport/time filters.")
    bullet(doc, "Phase 2 commerce module for SCL Direct subscriptions, premium access control, and subscription management.")
    bullet(doc, "QA, launch support, migration planning, analytics setup, and post-launch stabilization.")

    heading(doc, "Success Metrics")
    bullet(doc, "Users can create accounts, follow cappers, and find relevant cappers faster by sport, period, ROI, units, and win percentage.")
    bullet(doc, "All active Winible package links preserve SCL-controlled attribution and can be audited by admins.")
    bullet(doc, "Cappers can submit structured plays and package requests without bypassing SCL governance.")
    bullet(doc, "Admins can review capper activity, package changes, users, plays, and commerce status from a single dashboard.")
    bullet(doc, "The architecture can add SCL Direct subscriptions without rebuilding package, user, or access-control foundations.")

    heading(doc, "Assumptions and Open Questions")
    bullet(doc, "Existing capper, play, package, and leaderboard data can be exported or mapped from the current SCL system.")
    bullet(doc, "SCL will confirm current affiliate terms with Winible, including package-level attribution format and reporting cadence.")
    bullet(doc, "SCL will decide whether Phase 2 direct payments should make SCL or each capper the merchant of record.")
    bullet(doc, "SCL will define moderation rules, refund policy, age/geo restrictions, and responsible-gaming language before launch.")
    bullet(doc, "Final timeline and budget should follow discovery because the current backend, data quality, and migration needs are not yet confirmed.")

    heading(doc, "Suggested Next Step")
    paragraph(doc, "Approve a paid discovery/specification phase. The goal of discovery is to turn this proposal into an implementation-ready scope: sitemap, user roles, data model, package approval workflow, Stripe/Winible integration plan, migration plan, launch phases, timeline, and fixed build estimate.")

    heading(doc, "Source Notes")
    add_source(doc, "Sports Cappers Leaderboard public site", "https://www.sportscappersleaderboard.com/")
    add_source(doc, "Winible public site", "https://www.winible.com/")
    add_source(doc, "Winible Help Center: Onboarding-at-Home", "https://intercom.help/winible/en/articles/8895717-onboarding-at-home")
    add_source(doc, "Stripe Connect subscriptions documentation", "https://docs.stripe.com/connect/subscriptions")
    add_source(doc, "Stripe customer portal documentation", "https://docs.stripe.com/customer-management")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
